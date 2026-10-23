"""
CODEMEET 2026 - Comprehensive Panel Allocation & Dispatch Engine
================================================================
Parses all 57 teams across Panels A, B, C, D, E from raw data,
creates formatted Excel spreadsheet, and dispatches customized
Round 1 panel allocation emails to both Team Leaders and Team Members.
"""

import os
import sys
import re
import time
import argparse
import email.utils
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from typing import Dict, Any, List, Optional, Tuple
from dotenv import load_dotenv

# Ensure current dir is in sys.path
CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, CURRENT_DIR)

from email_service import get_smtp_config, _dispatch_smtp_message

load_dotenv()


def sanitize_email(raw_email: str) -> str:
    """Clean and correct common user typos in student emails."""
    em = (raw_email or "").strip().lower()
    if em.endswith("@gmail"):
        em += ".com"
    elif em.endswith("@gmailmcom"):
        em = em.replace("@gmailmcom", "@gmail.com")
    elif em.endswith("@gmial.com"):
        em = em.replace("@gmial.com", "@gmail.com")
    return em


def parse_raw_allocations(file_path: str) -> List[Dict[str, Any]]:
    """Parse the raw text file containing all panels (A, B, C, D, E) and 57 teams."""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Allocations file not found: {file_path}")

    with open(file_path, "r", encoding="utf-8") as f:
        text = f.read()

    panel_blocks = re.split(r"={10,}\s*\n(Panel [A-E][^\n]*)\n\s*={10,}", text)
    teams: List[Dict[str, Any]] = []

    for i in range(1, len(panel_blocks), 2):
        header = panel_blocks[i].strip()
        m = re.search(r"Panel ([A-E])", header)
        if not m:
            continue
        panel_letter = m.group(1).upper()
        panel_name = f"PANEL {panel_letter}"

        body = panel_blocks[i + 1].strip()
        team_chunks = re.split(r"\n(?=Team:)", body)

        for chunk in team_chunks:
            lines = [l.strip() for l in chunk.strip().split("\n") if l.strip()]
            entry: Dict[str, str] = {}
            for l in lines:
                if ":" in l:
                    k, v = l.split(":", 1)
                    entry[k.strip()] = v.strip()

            team_name = entry.get("Team", "Unknown Team")
            university = entry.get("University", "Participating Institution")
            leader_name = entry.get("Team Leader", "Team Leader")
            leader_email = sanitize_email(entry.get("Leader Email", ""))

            # Members list
            members: List[Dict[str, str]] = []
            # Leader is member 1
            if leader_email:
                members.append({
                    "name": leader_name,
                    "email": leader_email,
                    "role": "TEAM LEADER"
                })

            for num_word, label in [("Two", "MEMBER 2"), ("Three", "MEMBER 3"), ("Four", "MEMBER 4")]:
                m_name = entry.get(f"Member {num_word} Name", "").strip()
                m_email = sanitize_email(entry.get(f"Member {num_word} Email", ""))
                if m_email:
                    members.append({
                        "name": m_name or f"Member ({num_word})",
                        "email": m_email,
                        "role": label
                    })

            date_str = entry.get("Date", "23 October 2026")
            time_str = entry.get("Time", "As scheduled")
            wa_link = entry.get("WhatsApp", "")

            teams.append({
                "panel": panel_name,
                "panel_letter": panel_letter,
                "team_name": team_name,
                "university": university,
                "leader_name": leader_name,
                "leader_email": leader_email,
                "date": date_str,
                "time": time_str,
                "whatsapp_link": wa_link,
                "members": members
            })

    return teams


def export_allocations_to_excel(teams: List[Dict[str, Any]], output_path: str):
    """Generate a clean, styled Excel spreadsheet from parsed teams."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Round 1 Allocations"

    # Header styling
    header_fill = PatternFill(start_color="0A0A0F", end_color="0A0A0F", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="CCFF00")
    border_side = Side(border_style="thin", color="333333")
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    headers = [
        "Panel", "Team Name", "University", "Interview Date", "Time Slot",
        "Team Leader", "Leader Email", "Member 2 Name", "Member 2 Email",
        "Member 3 Name", "Member 3 Email", "Member 4 Name", "Member 4 Email",
        "WhatsApp Group Link"
    ]

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = cell_border
    ws.row_dimensions[1].height = 28

    panel_colors = {
        "PANEL A": "F4FFF0",
        "PANEL B": "F0F9FF",
        "PANEL C": "FAF5FF",
        "PANEL D": "FFFBEB",
        "PANEL E": "FFF1F2",
    }

    for row_idx, t in enumerate(teams, start=2):
        m2 = next((m for m in t["members"] if m["role"] == "MEMBER 2"), {})
        m3 = next((m for m in t["members"] if m["role"] == "MEMBER 3"), {})
        m4 = next((m for m in t["members"] if m["role"] == "MEMBER 4"), {})

        row_data = [
            t["panel"],
            t["team_name"],
            t["university"],
            t["date"],
            t["time"],
            t["leader_name"],
            t["leader_email"],
            m2.get("name", ""),
            m2.get("email", ""),
            m3.get("name", ""),
            m3.get("email", ""),
            m4.get("name", ""),
            m4.get("email", ""),
            t["whatsapp_link"]
        ]
        ws.append(row_data)

        fill_color = panel_colors.get(t["panel"], "FFFFFF")
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(name="Segoe UI", size=10)
            cell.alignment = center_align if col_idx in (1, 4, 5) else left_align
            cell.border = cell_border
            cell.fill = row_fill
        ws.row_dimensions[row_idx].height = 22

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    wb.save(output_path)
    print(f"Exported Excel spreadsheet to: {output_path}")


def generate_round1_email_html(
    team_name: str,
    university: str,
    panel: str,
    whatsapp_link: str,
    recipient_name: str,
    recipient_role: str,
    date_str: str,
    time_str: str,
    all_members: List[Dict[str, str]],
) -> str:
    """Generate sleek, responsive Cyberpunk HTML email for Round 1 Allocation."""

    roster_rows = ""
    for m in all_members:
        is_recip = m["email"].lower() == recipient_name.lower() or m["name"].lower() == recipient_name.lower()
        badge_style = "background-color: rgba(204, 255, 0, 0.15); border: 1px solid #ccff00; color: #ccff00;" if "LEADER" in m["role"] else "background-color: rgba(255, 255, 255, 0.08); border: 1px solid rgba(255, 255, 255, 0.15); color: #a1a1aa;"
        roster_rows += f"""
        <tr>
          <td style="padding: 8px 0; border-bottom: 1px solid rgba(255, 255, 255, 0.05);">
            <table border="0" cellpadding="0" cellspacing="0" width="100%">
              <tr>
                <td>
                  <span style="font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; font-size: 13px; font-weight: 600; color: #ffffff;">
                    {m['name']}
                  </span>
                  <span style="display: inline-block; margin-left: 8px; font-family: 'Courier New', Courier, monospace; font-size: 10px; font-weight: bold; padding: 2px 8px; border-radius: 4px; {badge_style}">
                    {m['role']}
                  </span>
                  <div style="font-size: 11px; color: #71717a; padding-top: 2px;">
                    {m['email']}
                  </div>
                </td>
              </tr>
            </table>
          </td>
        </tr>
        """

    return f"""<!DOCTYPE html PUBLIC "-//W3C//DTD XHTML 1.0 Transitional//EN" "http://www.w3.org/TR/xhtml1/DTD/xhtml1-transitional.dtd">
<html xmlns="http://www.w3.org/1999/xhtml" lang="en">
<head>
  <meta http-equiv="Content-Type" content="text/html; charset=UTF-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1.0" />
  <meta name="color-scheme" content="light dark" />
  <meta name="supported-color-schemes" content="light dark" />
  <title>CODEMEET 2026 - Round 1 Panel Allocation</title>
  <style type="text/css">
    body, table, td, p, a, li, blockquote {{ -webkit-text-size-adjust: 100%; -ms-text-size-adjust: 100%; }}
    table, td {{ mso-table-lspace: 0pt; mso-table-rspace: 0pt; }}
    img {{ -ms-interpolation-mode: bicubic; border: 0; outline: none; text-decoration: none; }}
    body {{ margin: 0 !important; padding: 0 !important; width: 100% !important; background-color: #050507; font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif; }}
    @media only screen and (max-width: 600px) {{
      .email-container {{ width: 100% !important; padding: 12px !important; }}
      .header-title {{ font-size: 26px !important; }}
      .panel-badge {{ font-size: 20px !important; }}
      .card-padding {{ padding: 20px 16px !important; }}
    }}
  </style>
</head>
<body style="margin: 0; padding: 32px 12px; background-color: #050507; color: #ffffff;">
  <center>
    <table class="email-container" border="0" cellpadding="0" cellspacing="0" width="100%" style="max-width: 580px; margin: 0 auto;">
      
      <!-- Top Cyber Tag Badge -->
      <tr>
        <td align="center" style="padding-bottom: 20px;">
          <table border="0" cellpadding="0" cellspacing="0">
            <tr>
              <td align="center" style="background-color: rgba(204, 255, 0, 0.08); border: 1px solid rgba(204, 255, 0, 0.4); border-radius: 24px; padding: 6px 16px;">
                <span style="font-family: 'Courier New', Courier, monospace; font-size: 11px; font-weight: bold; letter-spacing: 2px; color: #ccff00; text-transform: uppercase;">
                  // ROUND 1 ALLOCATION PROTOCOL
                </span>
              </td>
            </tr>
          </table>
        </td>
      </tr>

      <!-- Main Container Card -->
      <tr>
        <td>
          <table border="0" cellpadding="0" cellspacing="0" width="100%" style="background-color: #0a0a0f; border: 1px solid rgba(255, 255, 255, 0.12); border-radius: 16px; overflow: hidden; box-shadow: 0 12px 45px rgba(0, 0, 0, 0.85);">
            
            <!-- Glowing Neon Gradient Top Line -->
            <tr>
              <td height="4" style="background: linear-gradient(90deg, #ccff00 0%, #00f0ff 50%, #ccff00 100%);"></td>
            </tr>

            <!-- Header Title Block -->
            <tr>
              <td class="card-padding" style="padding: 32px 28px 18px 28px; text-align: center;">
                <div style="font-family: 'Courier New', Courier, monospace; font-size: 11px; letter-spacing: 3px; color: #a1a1aa; text-transform: uppercase; margin-bottom: 8px;">
                  SUIET &bull; WEBFLOW STUDENT COMMUNITY &bull; POWERED BY IBM
                </div>
                <h1 class="header-title" style="margin: 0; font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; font-size: 28px; font-weight: 900; letter-spacing: 1.5px; color: #ffffff; text-transform: uppercase;">
                  CODE<span style="color: #ccff00;">MEET</span> 2026
                </h1>
                <div style="margin-top: 6px; font-size: 13px; font-weight: 600; color: #00f0ff; letter-spacing: 0.5px;">
                  24H National Hackathon &mdash; Round 1 Online Screening
                </div>
              </td>
            </tr>

            <!-- Personalized Greeting Text -->
            <tr>
              <td class="card-padding" style="padding: 0 28px 20px 28px;">
                <p style="margin: 0; font-size: 14px; line-height: 22px; color: #e4e4e7;">
                  Greetings <strong>{recipient_name}</strong> ({recipient_role}) &amp; Team <strong>{team_name}</strong>,
                </p>
                <p style="margin: 10px 0 0 0; font-size: 13px; line-height: 21px; color: #a1a1aa;">
                  Congratulations on advancing to <strong>Round 1 (Online Screening)</strong> of CodeMeet 2026! Your team has been successfully allocated to <strong>{panel}</strong>. Below are your official evaluation slot details.
                </p>
              </td>
            </tr>

            <!-- Panel Allocation Highlight Card -->
            <tr>
              <td class="card-padding" style="padding: 0 28px 24px 28px;">
                <table border="0" cellpadding="0" cellspacing="0" width="100%" style="background-color: #111118; border: 1px solid rgba(204, 255, 0, 0.35); border-radius: 12px; padding: 22px; box-shadow: inset 0 0 20px rgba(204, 255, 0, 0.03);">
                  <tr>
                    <td>
                      <!-- Team & University Row -->
                      <table border="0" cellpadding="0" cellspacing="0" width="100%" style="margin-bottom: 14px; border-bottom: 1px solid rgba(255, 255, 255, 0.07); padding-bottom: 12px;">
                        <tr>
                          <td style="font-family: 'Courier New', Courier, monospace; font-size: 11px; color: #71717a; text-transform: uppercase; letter-spacing: 1px;">
                            REGISTERED TEAM &amp; INSTITUTION
                          </td>
                        </tr>
                        <tr>
                          <td style="font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; font-size: 18px; font-weight: 700; color: #ffffff; padding-top: 4px;">
                            {team_name}
                          </td>
                        </tr>
                        <tr>
                          <td style="font-size: 12px; color: #a1a1aa; padding-top: 2px;">
                            {university}
                          </td>
                        </tr>
                      </table>

                      <!-- Allocated Panel Badge -->
                      <table border="0" cellpadding="0" cellspacing="0" width="100%" style="margin-bottom: 14px; border-bottom: 1px solid rgba(255, 255, 255, 0.07); padding-bottom: 12px;">
                        <tr>
                          <td style="font-family: 'Courier New', Courier, monospace; font-size: 11px; color: #71717a; text-transform: uppercase; letter-spacing: 1px;">
                            ALLOCATED EVALUATION PANEL
                          </td>
                        </tr>
                        <tr>
                          <td style="padding-top: 6px;">
                            <span class="panel-badge" style="display: inline-block; background-color: rgba(204, 255, 0, 0.15); border: 1.5px solid #ccff00; color: #ccff00; font-family: 'Courier New', Courier, monospace; font-size: 22px; font-weight: 900; letter-spacing: 2px; padding: 8px 20px; border-radius: 8px; box-shadow: 0 0 20px rgba(204, 255, 0, 0.2);">
                              {panel}
                            </span>
                          </td>
                        </tr>
                      </table>

                      <!-- Date & Interview Slot Info -->
                      <table border="0" cellpadding="0" cellspacing="0" width="100%">
                        <tr>
                          <td style="font-family: 'Courier New', Courier, monospace; font-size: 11px; color: #71717a; text-transform: uppercase; letter-spacing: 1px;">
                            ROUND 1 INTERVIEW DATE &amp; TIME SLOT
                          </td>
                        </tr>
                        <tr>
                          <td style="font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; font-size: 14px; font-weight: 700; color: #38bdf8; padding-top: 4px;">
                            {date_str} &bull; {time_str}
                          </td>
                        </tr>
                        <tr>
                          <td style="font-size: 11px; color: #71717a; padding-top: 2px;">
                            Interview duration: 15 minutes per team
                          </td>
                        </tr>
                      </table>

                    </td>
                  </tr>
                </table>
              </td>
            </tr>

            <!-- WhatsApp Group CTA Section -->
            <tr>
              <td class="card-padding" style="padding: 0 28px 24px 28px;">
                <table border="0" cellpadding="0" cellspacing="0" width="100%" style="background: linear-gradient(135deg, rgba(37, 211, 102, 0.15) 0%, rgba(18, 140, 126, 0.08) 100%); border: 1px solid rgba(37, 211, 102, 0.45); border-radius: 12px; padding: 22px; text-align: center;">
                  <tr>
                    <td>
                      <div style="font-family: 'Courier New', Courier, monospace; font-size: 11px; font-weight: bold; letter-spacing: 2px; color: #25D366; text-transform: uppercase; margin-bottom: 6px;">
                        MANDATORY ACTION FOR ALL MEMBERS
                      </div>
                      <div style="font-size: 16px; font-weight: 800; color: #ffffff; margin-bottom: 8px;">
                        Join Your {panel} WhatsApp Group
                      </div>
                      <p style="margin: 0 0 16px 0; font-size: 12px; line-height: 18px; color: #d4d4d8;">
                        Live Google Meet presentation links, live evaluator queues, and real-time announcements will be shared exclusively inside this WhatsApp group.
                      </p>
                      
                      <!-- Green CTA Button -->
                      <table border="0" cellpadding="0" cellspacing="0" style="margin: 0 auto;">
                        <tr>
                          <td align="center" style="background-color: #25D366; border-radius: 10px; box-shadow: 0 4px 20px rgba(37, 211, 102, 0.4);">
                            <a href="{whatsapp_link}" target="_blank" style="display: inline-block; padding: 13px 30px; font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; font-size: 14px; font-weight: 800; letter-spacing: 0.5px; color: #000000; text-decoration: none; text-transform: uppercase;">
                              &rarr; JOIN {panel} WHATSAPP GROUP
                            </a>
                          </td>
                        </tr>
                      </table>

                      <div style="margin-top: 14px; font-size: 11px; color: #71717a; word-break: break-all;">
                        Direct Link: <a href="{whatsapp_link}" style="color: #25D366; text-decoration: underline;">{whatsapp_link}</a>
                      </div>
                    </td>
                  </tr>
                </table>
              </td>
            </tr>

            <!-- Team Members Roster -->
            <tr>
              <td class="card-padding" style="padding: 0 28px 24px 28px;">
                <table border="0" cellpadding="0" cellspacing="0" width="100%" style="background-color: rgba(255, 255, 255, 0.02); border: 1px solid rgba(255, 255, 255, 0.07); border-radius: 10px; padding: 16px;">
                  <tr>
                    <td style="font-family: 'Courier New', Courier, monospace; font-size: 11px; color: #ccff00; text-transform: uppercase; letter-spacing: 1.5px; padding-bottom: 10px; border-bottom: 1px solid rgba(255, 255, 255, 0.07);">
                      // REGISTERED TEAM ROSTER ({len(all_members)} MEMBERS)
                    </td>
                  </tr>
                  {roster_rows}
                </table>
              </td>
            </tr>

            <!-- Support Note -->
            <tr>
              <td class="card-padding" style="padding: 0 28px 24px 28px;">
                <table border="0" cellpadding="0" cellspacing="0" width="100%" style="background-color: rgba(255, 255, 255, 0.03); border: 1px solid rgba(255, 255, 255, 0.08); border-radius: 8px; padding: 14px;">
                  <tr>
                    <td style="font-size: 12px; line-height: 18px; color: #a1a1aa;">
                      <strong style="color: #ffffff;">Need Assistance or Have Questions?</strong><br />
                      Please coordinate with your designated Panel Coordinator inside the WhatsApp group, or reply directly to this email.
                    </td>
                  </tr>
                </table>
              </td>
            </tr>

            <!-- Footer Signoff -->
            <tr>
              <td style="padding: 22px 28px; background-color: #07070b; border-top: 1px solid rgba(255, 255, 255, 0.08); text-align: center;">
                <div style="font-size: 12px; font-weight: 700; color: #ffffff; letter-spacing: 0.5px;">
                  Best of Luck &bull; CodeMeet 2026 Organizing Committee
                </div>
                <div style="margin-top: 4px; font-size: 11px; color: #71717a;">
                  Srinivas University Institute of Engineering and Technology (SUIET), Mukka, Mangaluru
                </div>
              </td>
            </tr>

          </table>
        </td>
      </tr>

      <!-- Sub-Footer -->
      <tr>
        <td align="center" style="padding-top: 20px;">
          <p style="margin: 0; font-size: 10px; line-height: 16px; color: #52525b; text-align: center; font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;">
            This email was dispatched to verified participants of CODEMEET 2026.<br />
            &copy; 2026 SUIET &bull; Webflow Student Community. All rights reserved.
          </p>
        </td>
      </tr>

    </table>
  </center>
</body>
</html>
"""


def generate_round1_email_text(
    team_name: str,
    university: str,
    panel: str,
    whatsapp_link: str,
    recipient_name: str,
    recipient_role: str,
    date_str: str,
    time_str: str,
    all_members: List[Dict[str, str]],
) -> str:
    """Generate clean plain text fallback."""
    roster_lines = "\n".join([f"- {m['name']} ({m['role']}) : {m['email']}" for m in all_members])

    return f"""===================================================================
CODEMEET 2026 - 24H NATIONAL HACKATHON
ROUND 1 ONLINE SCREENING & PANEL ALLOCATION
===================================================================

Greetings {recipient_name} ({recipient_role}) & Team {team_name},

Congratulations on advancing to Round 1 (Online Screening) of CodeMeet 2026!
Your team has been officially allocated to {panel}.

-------------------------------------------------------------------
YOUR ALLOCATION DETAILS:
-------------------------------------------------------------------
* Registered Team    : {team_name}
* Institution        : {university}
* Allocated Panel    : {panel}
* Interview Date     : {date_str}
* Assigned Slot      : {time_str} (15 minutes duration)

-------------------------------------------------------------------
MANDATORY ACTION: JOIN YOUR PANEL WHATSAPP GROUP
-------------------------------------------------------------------
Live Google Meet presentation links, live evaluator queues, and real-time
announcements will be shared exclusively inside your panel's WhatsApp group.

Join {panel} WhatsApp Group:
{whatsapp_link}

-------------------------------------------------------------------
REGISTERED TEAM ROSTER:
-------------------------------------------------------------------
{roster_lines}

-------------------------------------------------------------------
Need assistance or have queries?
Reach out to your panel coordinator inside the WhatsApp group,
or reply directly to this email.

Best regards,
CODEMEET 2026 Organizing Committee
Srinivas University Institute of Engineering and Technology (SUIET), Mukka, Mangaluru
In collaboration with Webflow Student Community | Powered by IBM
"""


def send_single_allocation_email(
    to_email: str,
    team_name: str,
    university: str,
    panel: str,
    whatsapp_link: str,
    recipient_name: str,
    recipient_role: str,
    date_str: str,
    time_str: str,
    all_members: List[Dict[str, str]],
    custom_subject: Optional[str] = None
) -> Tuple[bool, str]:
    """Send an allocation email to a single recipient."""
    smtp_email, smtp_password, smtp_host, smtp_port = get_smtp_config()

    if not smtp_email or not smtp_password:
        return False, "SMTP credentials missing in .env"

    subject = custom_subject or f"[CODEMEET 2026] Round 1 Allocation: {panel} - Team {team_name}"

    try:
        msg = MIMEMultipart("alternative")
        msg["Subject"] = subject
        msg["From"] = f"CODEMEET 2026 Hackathon Desk <{smtp_email}>"
        msg["To"] = to_email
        msg["Reply-To"] = smtp_email
        msg["Date"] = email.utils.formatdate(localtime=True)
        msg["Message-ID"] = email.utils.make_msgid(domain="srinivasuniversity.edu.in")
        msg["X-Priority"] = "1"
        msg["Importance"] = "High"

        plain_text = generate_round1_email_text(
            team_name=team_name,
            university=university,
            panel=panel,
            whatsapp_link=whatsapp_link,
            recipient_name=recipient_name,
            recipient_role=recipient_role,
            date_str=date_str,
            time_str=time_str,
            all_members=all_members
        )
        html_content = generate_round1_email_html(
            team_name=team_name,
            university=university,
            panel=panel,
            whatsapp_link=whatsapp_link,
            recipient_name=recipient_name,
            recipient_role=recipient_role,
            date_str=date_str,
            time_str=time_str,
            all_members=all_members
        )

        part1 = MIMEText(plain_text, "plain", "utf-8")
        part2 = MIMEText(html_content, "html", "utf-8")
        msg.attach(part1)
        msg.attach(part2)

        return _dispatch_smtp_message(to_email, msg)
    except Exception as e:
        return False, f"Failed to build/dispatch: {str(e)}"


def send_panel_samples_to_admin(admin_email: str, raw_file_path: str):
    """
    Send 5 sample emails (one for Panel A, B, C, D, E) to the admin email.
    Allows admin to inspect the formatting and links for each panel.
    """
    teams = parse_raw_allocations(raw_file_path)
    panels = ["PANEL A", "PANEL B", "PANEL C", "PANEL D", "PANEL E"]

    print(f"\n=======================================================")
    print(f"DISPATCHING 5 PANEL SAMPLES TO ADMIN: {admin_email}")
    print(f"=======================================================")

    for p in panels:
        # Find first team in this panel
        team_entry = next((t for t in teams if t["panel"] == p), None)
        if not team_entry:
            print(f"No team found for {p}")
            continue

        print(f"\nDispatching {p} sample [Team: {team_entry['team_name']} | Slot: {team_entry['time']}] -> {admin_email}...", end="", flush=True)

        ok, msg = send_single_allocation_email(
            to_email=admin_email,
            team_name=team_entry["team_name"],
            university=team_entry["university"],
            panel=team_entry["panel"],
            whatsapp_link=team_entry["whatsapp_link"],
            recipient_name="Shreyas Sondur",
            recipient_role="TEAM LEADER (PREVIEW)",
            date_str=team_entry["date"],
            time_str=team_entry["time"],
            all_members=team_entry["members"],
            custom_subject=f"[CODEMEET 2026] [SAMPLE {p}] Round 1 Allocation - Team {team_entry['team_name']}"
        )

        if ok:
            print(" [OK]")
        else:
            print(f" [FAILED: {msg}]")

        time.sleep(2.0)  # Gentle delay between preview emails

    print("\n=======================================================")
    print(f"ALL 5 PANEL SAMPLES DISPATCHED TO: {admin_email}")
    print("=======================================================\n")


def dispatch_to_all_teams(raw_file_path: str, dry_run: bool = False, delay_seconds: float = 1.2):
    """Dispatch allocation emails to both Team Leaders and all Members across all 57 teams."""
    teams = parse_raw_allocations(raw_file_path)
    total_teams = len(teams)
    total_recipients = sum(len(t["members"]) for t in teams)

    print(f"\n=======================================================")
    print(f"CODE MEET 2026 - BATCH ALLOCATION DISPATCH")
    print(f"Total Teams: {total_teams} across Panels A, B, C, D, E")
    print(f"Total Individual Recipients (Leaders + Members): {total_recipients}")
    print(f"Dry Run Mode: {dry_run}")
    print(f"=======================================================\n")

    if dry_run:
        for idx, t in enumerate(teams, start=1):
            print(f"[{idx}/{total_teams}] {t['panel']} | Team '{t['team_name']}' | Slot: {t['time']} | Members: {len(t['members'])}")
            for m in t["members"]:
                print(f"    -> {m['name']} ({m['role']}) : {m['email']}")
        print("\n[DRY RUN FINISHED - No emails were sent]")
        return

    sent_count = 0
    failed_count = 0
    failures = []
    log_records = []
    log_file_path = os.path.join(CURRENT_DIR, "data", "dispatch_status.json")

    for t_idx, t in enumerate(teams, start=1):
        print(f"\n--- [{t_idx}/{total_teams}] Team '{t['team_name']}' ({t['panel']}) ---")
        for m in t["members"]:
            print(f"  Sending to {m['name']} ({m['role']}) <{m['email']}>...", end="", flush=True)
            ok, msg = send_single_allocation_email(
                to_email=m["email"],
                team_name=t["team_name"],
                university=t["university"],
                panel=t["panel"],
                whatsapp_link=t["whatsapp_link"],
                recipient_name=m["name"],
                recipient_role=m["role"],
                date_str=t["date"],
                time_str=t["time"],
                all_members=t["members"]
            )
            rec_status = {
                "timestamp": time.strftime("%Y-%m-%d %H:%M:%S"),
                "panel": t["panel"],
                "team": t["team_name"],
                "name": m["name"],
                "role": m["role"],
                "email": m["email"],
                "status": "SENT" if ok else "FAILED",
                "detail": msg
            }
            log_records.append(rec_status)

            if ok:
                print(" [OK]")
                sent_count += 1
            else:
                print(f" [FAILED: {msg}]")
                failed_count += 1
                failures.append({"team": t["team_name"], "recipient": m["email"], "error": msg})

            # Save progress periodically
            import json
            with open(log_file_path, "w", encoding="utf-8") as lf:
                json.dump({
                    "total_teams": total_teams,
                    "total_recipients": total_recipients,
                    "sent_count": sent_count,
                    "failed_count": failed_count,
                    "records": log_records
                }, lf, indent=2)

            if delay_seconds > 0:
                time.sleep(delay_seconds)

    print("\n=======================================================")
    print(f"BATCH COMPLETE! Sent: {sent_count} | Failed: {failed_count} out of {total_recipients}")
    print(f"Dispatch status saved to: {log_file_path}")
    if failures:
        print(f"Failures ({len(failures)}):")
        for f in failures:
            print(f"  - {f['team']} ({f['recipient']}): {f['error']}")
    print("=======================================================\n")


def main():
    parser = argparse.ArgumentParser(description="CodeMeet 2026 Panel Allocation Engine")
    parser.add_argument("--send-samples", type=str, help="Send 5 panel sample emails (Panels A-E) to specified admin email")
    parser.add_argument("--export-excel", action="store_true", help="Export allocations to an Excel file")
    parser.add_argument("--dispatch-all", action="store_true", help="Dispatch emails to all 57 teams (Leaders + Members)")
    parser.add_argument("--dry-run", action="store_true", help="Dry run mode without sending emails")
    parser.add_argument("--delay", type=float, default=1.2, help="Delay between emails in seconds (default: 1.2)")

    args = parser.parse_args()
    raw_path = os.path.join(CURRENT_DIR, "data", "panel_allocations_raw.txt")
    excel_path = os.path.join(CURRENT_DIR, "data", "panel_allocations.xlsx")

    if args.export_excel:
        teams = parse_raw_allocations(raw_path)
        export_allocations_to_excel(teams, excel_path)
        return

    if args.send_samples:
        send_panel_samples_to_admin(args.send_samples.strip(), raw_path)
        return

    if args.dispatch_all:
        dispatch_to_all_teams(raw_path, dry_run=args.dry_run, delay_seconds=args.delay)
        return

    print("Please specify an action:\n  --send-samples <email>\n  --export-excel\n  --dispatch-all [--dry-run]")


if __name__ == "__main__":
    main()
